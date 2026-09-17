# -*- coding: utf-8 -*-
"""
Экспорт актов:
1. build_act_docx() — один акт (из /actcalc, т.е. db.material_acts) в .docx
   по точному образцу шаблона «АВР_Монтаж_Кондиционеров»: шапка (номер,
   дата, модель, заказчик, адрес, исполнитель), таблица позиций (работы +
   материалы, с ценами и итогом), стандартный текст и подписи.
2. build_acts_summary_xlsx() — сводная таблица по ВСЕМ актам (и простые
   фото-акты из db.acts, и структурированные из db.material_acts) в один
   Excel-файл, для бухгалтерии/отчётности.

Шапка акта (заказчик/адрес/модель) в /actcalc не запрашивается — там
считаются только работы/материалы. Если админ передаёт эти данные при
экспорте (аргументами команды), они подставляются; если нет — остаются
пустой линией, как на бумажном бланке, который потом дозаполняют от руки.
"""

from datetime import datetime, date

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import act_catalog

FOOTER_TEXT = [
    "Работы выполнены в установленные сроки, в полном объеме и с надлежащим качеством.",
    "Место установки оборудования согласованно с заказчиком.",
    "Претензий друг к другу стороны не имеют.",
    "Инструкция по эксплуатации, пульт ДУ переданы заказчику.",
]

TABLE_HEADERS = ["№", "Наименование", "Ед. изм.", "Кол-во", "Цена", "Стоимость"]
COL_WIDTHS_CM = [1.0, 7.0, 2.0, 2.0, 2.5, 3.0]


def _shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 10, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def build_act_docx(
    act_id: int,
    created_at: str,
    items: list[dict],
    total: float,
    capacity_group: str,
    installer_name: str,
    output_path: str,
    client_name: str | None = None,
    address: str | None = None,
    ac_model: str | None = None,
) -> None:
    """items — список {"code":, "name":, "unit":, "qty":, "unit_price":, "subtotal":},
    как хранится в db.material_acts.items. Разбивает их на «Работы» и
    «Материалы» по act_catalog.ALL_ITEMS_BY_CODE (is_material)."""
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    try:
        created_date = datetime.fromisoformat(created_at).date()
    except ValueError:
        created_date = date.today()
    date_str = f"{created_date.day:02d} {_RU_MONTHS[created_date.month]} {created_date.year}"

    title = doc.add_paragraph()
    title_run = title.add_run(f"Акт выполненных работ № {act_id}")
    title_run.bold = True
    title_run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.add_run(f"от {date_str} г.    Кондиционер – модель: ").font.size = Pt(11)
    p.add_run(ac_model or "_" * 40).font.size = Pt(11)

    p = doc.add_paragraph()
    p.add_run("Заказчик: ").font.size = Pt(11)
    p.add_run(client_name or "_" * 55).font.size = Pt(11)

    p = doc.add_paragraph()
    p.add_run("Адрес монтажа: ").font.size = Pt(11)
    p.add_run(address or "_" * 50).font.size = Pt(11)

    p = doc.add_paragraph()
    p.add_run("Исполнитель: ").font.size = Pt(11)
    p.add_run(installer_name or "_" * 40).font.size = Pt(11)

    p = doc.add_paragraph()
    p.add_run(f"Группа мощности: {capacity_group}").italic = True

    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate(COL_WIDTHS_CM):
        table.columns[i].width = Cm(w)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(TABLE_HEADERS):
        _set_cell_text(hdr_cells[i], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade_cell(hdr_cells[i], "D9E2F3")
        hdr_cells[i].width = Cm(COL_WIDTHS_CM[i])

    works = [it for it in items if not act_catalog.ALL_ITEMS_BY_CODE.get(it["code"], {}).get("is_material", True)]
    materials = [it for it in items if act_catalog.ALL_ITEMS_BY_CODE.get(it["code"], {}).get("is_material", True)]

    def add_section_row(label: str) -> None:
        row = table.add_row()
        row.cells[0].merge(row.cells[-1])
        _set_cell_text(row.cells[0], label, bold=True, size=10)
        _shade_cell(row.cells[0], "F2F2F2")

    def add_item_row(idx: int, it: dict) -> None:
        row = table.add_row().cells
        _set_cell_text(row[0], str(idx), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(row[1], it["name"], size=10)
        _set_cell_text(row[2], it["unit"], size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(row[3], f"{it['qty']:g}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(row[4], f"{it['unit_price']:.2f}", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _set_cell_text(row[5], f"{it['subtotal']:.2f}", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    idx = 1
    if works:
        add_section_row("Работы")
        for it in works:
            add_item_row(idx, it)
            idx += 1
    if materials:
        add_section_row("Материалы")
        for it in materials:
            add_item_row(idx, it)
            idx += 1

    total_row = table.add_row()
    total_row.cells[0].merge(total_row.cells[-2])
    _set_cell_text(total_row.cells[0], "Итого к оплате:", bold=True, size=11)
    _set_cell_text(total_row.cells[-1], f"{total:.2f}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _shade_cell(total_row.cells[0], "FFF2CC")
    _shade_cell(total_row.cells[-1], "FFF2CC")

    doc.add_paragraph()
    for line in FOOTER_TEXT:
        doc.add_paragraph(line)

    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.add_run("Исполнитель: ").bold = True
    sig.add_run("_" * 20 + f"  {installer_name or ''}" + " " * 10)
    sig.add_run("Заказчик: ").bold = True
    sig.add_run("_" * 25)

    doc.save(output_path)


_RU_MONTHS = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля", 5: "мая", 6: "июня",
    7: "июля", 8: "августа", 9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
}


def build_acts_summary_xlsx(
    photo_acts: list[dict],
    material_acts: list[dict],
    output_path: str,
) -> None:
    """Сводная таблица по ВСЕМ актам — и простым (фото+адрес), и
    структурированным (/actcalc, с ценами). Один лист, одна строка на акт."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Акты"
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["ID", "Дата", "Тип", "Адрес/Заказчик", "Группа/Модель", "Сумма", "Исполнитель", "Статус"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = border

    row_idx = 2
    combined = []
    for a in photo_acts:
        combined.append({
            "id": a["id"], "created_at": a["created_at"], "type": "Фото",
            "who_where": a["address"], "detail": "", "amount": None,
            "installer": a.get("installer_chat_id"), "status": a.get("status", ""),
        })
    for a in material_acts:
        combined.append({
            "id": a["id"], "created_at": a["created_at"], "type": "Расчёт (/actcalc)",
            "who_where": "", "detail": a["capacity_group"], "amount": a["total"],
            "installer": a.get("installer_chat_id"), "status": "",
        })
    combined.sort(key=lambda r: r["created_at"], reverse=True)

    for r in combined:
        try:
            dt_str = datetime.fromisoformat(r["created_at"]).strftime("%d.%m.%Y %H:%M")
        except ValueError:
            dt_str = r["created_at"]
        values = [
            r["id"], dt_str, r["type"], r["who_where"], r["detail"],
            r["amount"], r["installer"], r["status"],
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col, value=v)
            c.border = border
            if col == 6 and v is not None:
                c.number_format = "0.00"
        row_idx += 1

    widths = [8, 16, 16, 26, 16, 12, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_path)
